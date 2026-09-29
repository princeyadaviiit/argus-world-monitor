import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Exact Reference Image Color Palette
C_NAVY         = RGBColor(27, 54, 93)     # #1B365D - Deep Navy (PUNK RECORDS / Table Header)
C_TEAL_SLATE   = RGBColor(27, 73, 101)    # #1B4965 - Deep Teal/Slate (Citizen Vault)
C_FOREST_GREEN = RGBColor(30, 126, 52)    # #1E7E34 - Forest Green (TRAFFIC SATELLITE / OCR)
C_BURNT_ORANGE = RGBColor(192, 86, 33)    # #C05621 - Burnt Orange (BANKING SATELLITE / Checksum)
C_ROYAL_PURPLE = RGBColor(90, 55, 140)    # #5A378C - Deep Purple (LEGAL SATELLITE / Output)
C_MAROON       = RGBColor(140, 45, 25)    # #8C2D19 - Deep Maroon/Rust (CIVIC LITERACY / Crisis)
C_DARK_MAROON  = RGBColor(128, 24, 24)    # #801818 - Dark Crimson Banner (Why this matters)
C_WHITE        = RGBColor(255, 255, 255)  # #FFFFFF - Crisp pure white
C_OFFWHITE     = RGBColor(248, 250, 252)  # #F8FAFC - Soft Slate
C_LIGHT_BLUE   = RGBColor(239, 246, 255)  # #EFF6FF - Soft Ice Blue
C_TEXT_DARK    = RGBColor(15, 23, 42)     # #0F172A - Deep Charcoal
C_TEXT_MUTED   = RGBColor(71, 85, 105)    # #475569 - Muted Slate
C_BORDER_LIGHT = RGBColor(203, 213, 225)  # #CBD5E1 - Light Slate Border

prs = pptx.Presentation('SIH26163_ARGUS_Idea_PPT_Color.pptx')

def apply_solid_fill(shape, color, border_color=None, border_width=Pt(1)):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = border_width
    else:
        shape.line.fill.background()

def set_tf_color(tf, color, bold=None, size=None):
    for p in tf.paragraphs:
        if bold is not None:
            p.font.bold = bold
        if size is not None:
            p.font.size = size
        for run in p.runs:
            run.font.color.rgb = color
            if bold is not None:
                run.font.bold = bold
            if size is not None:
                run.font.size = size

# ==============================================================================
# SLIDE 1: Title Slide
# ==========================================
s1 = prs.slides[0]
for s in s1.shapes:
    if s.name == 'Shape 0':
        apply_solid_fill(s, C_NAVY)
    elif s.name == 'Shape 13':  # Idea title box background -> Solid Deep Navy Box!
        apply_solid_fill(s, C_NAVY, border_color=C_TEAL_SLATE, border_width=Pt(1.5))
    elif s.name == 'Text 14':  # Idea title text
        for p in s.text_frame.paragraphs:
            for r in p.runs:
                if 'Idea Title' in r.text or 'ARGUS' in r.text:
                    r.font.color.rgb = C_WHITE
                    r.font.bold = True
                else:
                    r.font.color.rgb = RGBColor(219, 234, 254)
    elif s.name in ['Shape 3', 'Shape 5', 'Shape 7', 'Shape 9', 'Shape 11']:
        dots = {'Shape 3': C_FOREST_GREEN, 'Shape 5': C_BURNT_ORANGE, 'Shape 7': C_ROYAL_PURPLE, 'Shape 9': C_NAVY, 'Shape 11': C_MAROON}
        apply_solid_fill(s, dots[s.name])

# ==============================================================================
# SLIDE 2: Proposed Solution
# ==============================================================================
s2 = prs.slides[1]

for s in s2.shapes:
    # Oval team logo at top left
    if s.name == 'Shape 0':
        apply_solid_fill(s, C_WHITE, border_color=C_ROYAL_PURPLE, border_width=Pt(1.5))
        if s.has_text_frame:
            set_tf_color(s.text_frame, C_TEXT_DARK, bold=True)

    # Scope Container (Shape 9)
    elif s.name == 'Shape 9':
        apply_solid_fill(s, C_OFFWHITE, border_color=C_NAVY, border_width=Pt(1.5))
    
    # 4 Solution Flow Boxes: Identify, Correlate, Prove, Remediate
    # Solid background boxes with crisp white text!
    elif s.name == 'Shape 12': # 1. Identify: Deep Teal
        apply_solid_fill(s, C_TEAL_SLATE)
    elif s.name == 'Shape 17': # 2. Correlate: Burnt Orange
        apply_solid_fill(s, C_BURNT_ORANGE)
    elif s.name == 'Shape 22': # 3. Prove: Forest Green
        apply_solid_fill(s, C_FOREST_GREEN)
    elif s.name == 'Shape 27': # 4. Remediate: Royal Purple
        apply_solid_fill(s, C_ROYAL_PURPLE)

    # Circle icon badges inside the flow boxes
    elif s.name == 'Shape 13':
        apply_solid_fill(s, RGBColor(20, 55, 76))
    elif s.name == 'Shape 18':
        apply_solid_fill(s, RGBColor(156, 68, 24))
    elif s.name == 'Shape 23':
        apply_solid_fill(s, RGBColor(22, 94, 38))
    elif s.name == 'Shape 28':
        apply_solid_fill(s, RGBColor(70, 42, 110))

    # Text for 4 Solution Flow Boxes -> Bold White titles, White subtitles!
    elif s.name in ['Text 14', 'Text 19', 'Text 24', 'Text 29']:
        set_tf_color(s.text_frame, C_WHITE, bold=True)
    elif s.name in ['Text 15', 'Text 20', 'Text 25', 'Text 30']:
        set_tf_color(s.text_frame, C_WHITE)

    # Connecting arrows
    elif s.name in ['Shape 16', 'Shape 21', 'Shape 26']:
        try:
            s.line.color.rgb = C_NAVY
            s.line.width = Pt(2)
        except Exception:
            pass

    # 4 Innovation & Rigor Cards on Right:
    # Give them solid semantic colored cards with white text, or clean themed cards!
    # In the reference image, the satellite boxes are solid rich colors.
    # Giving each Innovation card a distinct themed border & background:
    elif s.name == 'Shape 32': # Two-Signal: Deep Teal/Navy
        apply_solid_fill(s, C_OFFWHITE, border_color=C_TEAL_SLATE, border_width=Pt(1.5))
    elif s.name == 'Shape 34': # Attack-Surface: Forest Green
        apply_solid_fill(s, C_OFFWHITE, border_color=C_FOREST_GREEN, border_width=Pt(1.5))
    elif s.name == 'Shape 36': # Algorithmic Scoring: Burnt Orange
        apply_solid_fill(s, C_OFFWHITE, border_color=C_BURNT_ORANGE, border_width=Pt(1.5))
    elif s.name == 'Shape 38': # OWASP LLM Suite: Royal Purple
        apply_solid_fill(s, C_OFFWHITE, border_color=C_ROYAL_PURPLE, border_width=Pt(1.5))

    # Real Evidenced Vulnerabilities Table:
    elif s.has_table:
        table = s.table
        # Header Row: Solid Deep Navy Blue with White Text (Reference Image table style)
        for cell in table.rows[0].cells:
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_NAVY
            for p in cell.text_frame.paragraphs:
                p.font.bold = True
                for run in p.runs:
                    run.font.color.rgb = C_WHITE
        # Alternating Rows
        for r_idx in range(1, len(table.rows)):
            row = table.rows[r_idx]
            bg_col = C_WHITE if r_idx % 2 == 1 else C_OFFWHITE
            for cell in row.cells:
                cell.fill.solid()
                cell.fill.fore_color.rgb = bg_col

# ==============================================================================
# SLIDE 3: Technical Approach
# ==============================================================================
s3 = prs.slides[2]

# Top 4 Architecture Diagram boxes: Solid semantic colors matching Reference Image 2!
for s in s3.shapes:
    if s.name == 'Shape 7':   # Target Platform: Deep Teal
        apply_solid_fill(s, C_TEAL_SLATE)
    elif s.name == 'Shape 12': # Multi-Engine Scanners: Forest Green
        apply_solid_fill(s, C_FOREST_GREEN)
    elif s.name == 'Shape 17': # Correlation Engine: Burnt Orange
        apply_solid_fill(s, C_BURNT_ORANGE)
    elif s.name == 'Shape 22': # Outputs & Evidence: Royal Purple
        apply_solid_fill(s, C_ROYAL_PURPLE)

    # Circle icon backgrounds
    elif s.name == 'Shape 8':
        apply_solid_fill(s, RGBColor(20, 55, 76))
    elif s.name == 'Shape 13':
        apply_solid_fill(s, RGBColor(22, 94, 38))
    elif s.name == 'Shape 18':
        apply_solid_fill(s, RGBColor(156, 68, 24))
    elif s.name == 'Shape 23':
        apply_solid_fill(s, RGBColor(70, 42, 110))

    # Architecture text -> Bold White titles, White bullet points!
    elif s.name in ['Text 9', 'Text 14', 'Text 19', 'Text 24']:
        set_tf_color(s.text_frame, C_WHITE, bold=True)
    elif s.name in ['Text 10', 'Text 15', 'Text 20', 'Text 25']:
        set_tf_color(s.text_frame, C_WHITE)

    # Connecting horizontal arrows between top boxes
    elif s.name in ['Shape 11', 'Shape 16', 'Shape 21']:
        try:
            s.line.color.rgb = C_NAVY
            s.line.width = Pt(2)
        except Exception:
            pass

    # 8 Flowchart Steps: Solid vibrant boxes matching Reference Image 2!
    elif s.name == 'Shape 28': # 1 Scope Gate: Deep Navy
        apply_solid_fill(s, C_NAVY)
    elif s.name == 'Shape 31': # 2 Recon: Deep Teal
        apply_solid_fill(s, C_TEAL_SLATE)
    elif s.name == 'Shape 34': # 3 Multi-Scan: Forest Green
        apply_solid_fill(s, C_FOREST_GREEN)
    elif s.name == 'Shape 37': # 4 Correlate: Burnt Orange
        apply_solid_fill(s, C_BURNT_ORANGE)
    elif s.name == 'Shape 44': # 5 Confirm: Deep Maroon
        apply_solid_fill(s, C_MAROON)
    elif s.name == 'Shape 47': # 6 Prove: Forest Green
        apply_solid_fill(s, C_FOREST_GREEN)
    elif s.name == 'Shape 50': # 7 Score: Burnt Orange
        apply_solid_fill(s, C_BURNT_ORANGE)
    elif s.name == 'Shape 53': # 8 Report: Royal Purple
        apply_solid_fill(s, C_ROYAL_PURPLE)

    # Flowchart text -> Bold White titles, White subtitles!
    elif s.name in ['Text 29', 'Text 32', 'Text 35', 'Text 38', 'Text 45', 'Text 48', 'Text 51', 'Text 54']:
        set_tf_color(s.text_frame, C_WHITE, bold=True)
    elif s.name in ['Text 30', 'Text 33', 'Text 36', 'Text 39', 'Text 46', 'Text 49', 'Text 52', 'Text 55']:
        set_tf_color(s.text_frame, C_WHITE)

    # Connectors between flowchart steps
    elif s.name in ['Shape 40', 'Shape 41', 'Shape 42', 'Shape 43', 'Shape 56', 'Shape 57', 'Shape 58']:
        try:
            s.line.color.rgb = C_NAVY
            s.line.width = Pt(1.5)
        except Exception:
            pass

    # Right Box: Technical Architecture & Specifications
    elif s.name == 'Shape 61':
        apply_solid_fill(s, C_OFFWHITE, border_color=C_NAVY, border_width=Pt(1.5))

# ==============================================================================
# SLIDE 4: Feasibility and Viability (Redesigned Flowchart + Explanations)
# ==============================================================================
s4 = prs.slides[3]

# Left Column: 5 Process Flow Nodes (Solid Background Boxes from Reference Image 2!)
flow_nodes_s4 = {
    'Rounded Rectangle 10': C_TEAL_SLATE,   # 1. Sandbox: Deep Teal
    'Rounded Rectangle 12': C_FOREST_GREEN, # 2. Surface Recon: Forest Green
    'Rounded Rectangle 14': C_BURNT_ORANGE, # 3. Two-Signal Gate: Burnt Orange
    'Rounded Rectangle 16': C_ROYAL_PURPLE, # 4. Safe PoC: Royal Purple
    'Rounded Rectangle 18': C_MAROON        # 5. Hardened Fixes: Deep Maroon
}

# Right Column: 4 Explanation Cards (Clean containers with themed borders)
expl_cards_s4 = {
    'Rounded Rectangle 19': (C_OFFWHITE, C_TEAL_SLATE),   # Technical Feasibility
    'Rounded Rectangle 20': (C_OFFWHITE, C_FOREST_GREEN), # Operational & Economic Viability
    'Rounded Rectangle 21': (C_OFFWHITE, C_BURNT_ORANGE), # Challenge & Risk Mitigation
    'Rounded Rectangle 22': (C_OFFWHITE, C_ROYAL_PURPLE)  # Strategic Impact & Scalability
}

for s in s4.shapes:
    if s.name in flow_nodes_s4:
        fill_col = flow_nodes_s4[s.name]
        apply_solid_fill(s, fill_col)
        # Text is crisp Pure White!
        for p in s.text_frame.paragraphs:
            p.font.color.rgb = C_WHITE
            for r in p.runs:
                r.font.color.rgb = C_WHITE
    elif s.name in expl_cards_s4:
        bg_col, border_col = expl_cards_s4[s.name]
        apply_solid_fill(s, bg_col, border_color=border_col, border_width=Pt(1.5))
        if s.text_frame and len(s.text_frame.paragraphs) > 0:
            p0 = s.text_frame.paragraphs[0]
            p0.font.bold = True
            for r in p0.runs:
                r.font.color.rgb = border_col
            for p in s.text_frame.paragraphs[1:]:
                for r in p.runs:
                    r.font.color.rgb = C_TEXT_DARK
    elif 'Arrow' in s.name or 'Shape' in s.name:
        try:
            s.line.color.rgb = C_NAVY
            s.line.width = Pt(2)
        except Exception:
            pass

# ==============================================================================
# SLIDE 5: Impact and Benefits
# ==============================================================================
s5 = prs.slides[4]

for s in s5.shapes:
    # "Who Benefits" Icon circle badges (1, 2, 3, 4)
    if s.name == 'Shape 7':   # NTRO
        apply_solid_fill(s, C_FOREST_GREEN)
    elif s.name == 'Shape 9':  # Developers
        apply_solid_fill(s, C_BURNT_ORANGE)
    elif s.name == 'Shape 11': # Analysts
        apply_solid_fill(s, C_ROYAL_PURPLE)
    elif s.name == 'Shape 13': # DevSecOps
        apply_solid_fill(s, C_NAVY)

    # "Benefits" Icon circle badges (5, 6)
    elif s.name == 'Shape 16': # National Defense
        apply_solid_fill(s, C_NAVY)
    elif s.name == 'Shape 18': # Zero Footprint
        apply_solid_fill(s, C_FOREST_GREEN)

    # Scalability Path Boxes: NOW, NEXT, FUTURE
    # Solid background boxes with Bold White text!
    elif s.name == 'Shape 21': # NOW: Solid Forest Green
        apply_solid_fill(s, C_FOREST_GREEN)
    elif s.name == 'Text 22':
        set_tf_color(s.text_frame, C_WHITE, bold=True)
    elif s.name == 'Text 23':
        set_tf_color(s.text_frame, C_WHITE, bold=True)
    elif s.name == 'Text 24':
        set_tf_color(s.text_frame, RGBColor(220, 252, 231)) # Light mint white

    elif s.name == 'Shape 26': # NEXT: Solid Burnt Orange
        apply_solid_fill(s, C_BURNT_ORANGE)
    elif s.name == 'Text 27':
        set_tf_color(s.text_frame, C_WHITE, bold=True)
    elif s.name == 'Text 28':
        set_tf_color(s.text_frame, C_WHITE, bold=True)
    elif s.name == 'Text 29':
        set_tf_color(s.text_frame, RGBColor(254, 243, 199)) # Light amber white

    elif s.name == 'Shape 31': # FUTURE: Solid Royal Purple
        apply_solid_fill(s, C_ROYAL_PURPLE)
    elif s.name == 'Text 32':
        set_tf_color(s.text_frame, C_WHITE, bold=True)
    elif s.name == 'Text 33':
        set_tf_color(s.text_frame, C_WHITE, bold=True)
    elif s.name == 'Text 34':
        set_tf_color(s.text_frame, RGBColor(243, 232, 255)) # Light lavender white

    # Connectors between NOW, NEXT, FUTURE
    elif s.name in ['Shape 25', 'Shape 30']:
        try:
            s.line.color.rgb = C_NAVY
            s.line.width = Pt(2)
        except Exception:
            pass

    # Deliverable per Vulnerability (8 badges: Shapes 36, 39, 42, 45, 48, 51, 54, 57)
    # Solid colored square badge boxes with Bold White numbers!
    elif s.name in ['Shape 36', 'Shape 48']: # 1, 5 -> Deep Navy
        apply_solid_fill(s, C_NAVY)
    elif s.name in ['Shape 39', 'Shape 51']: # 2, 6 -> Forest Green
        apply_solid_fill(s, C_FOREST_GREEN)
    elif s.name in ['Shape 42', 'Shape 54']: # 3, 7 -> Burnt Orange
        apply_solid_fill(s, C_BURNT_ORANGE)
    elif s.name in ['Shape 45', 'Shape 57']: # 4, 8 -> Royal Purple
        apply_solid_fill(s, C_ROYAL_PURPLE)
    elif s.name in ['Text 37', 'Text 40', 'Text 43', 'Text 46', 'Text 49', 'Text 52', 'Text 55', 'Text 58']:
        set_tf_color(s.text_frame, C_WHITE, bold=True)

# ==============================================================================
# SLIDE 6: Research and References
# ==============================================================================
s6 = prs.slides[5]

# 8 Number Badges: Solid colored boxes with Bold White numbers!
badge_colors_s6 = {
    'Shape 6':  C_NAVY,         # 1
    'Shape 11': C_FOREST_GREEN, # 2
    'Shape 16': C_BURNT_ORANGE, # 3
    'Shape 21': C_ROYAL_PURPLE, # 4
    'Shape 26': C_NAVY,         # 5
    'Shape 31': C_FOREST_GREEN, # 6
    'Shape 36': C_BURNT_ORANGE, # 7
    'Shape 41': C_ROYAL_PURPLE  # 8
}

text_num_s6 = ['Text 7', 'Text 12', 'Text 17', 'Text 22', 'Text 27', 'Text 32', 'Text 37', 'Text 42']

for s in s6.shapes:
    if s.name in badge_colors_s6:
        apply_solid_fill(s, badge_colors_s6[s.name])
    elif s.name in text_num_s6:
        set_tf_color(s.text_frame, C_WHITE, bold=True)
    elif s.name == 'Shape 46': # Bottom Ethical Scope container
        # Modeled directly after Reference Image 1 top banner ("WHY THIS MATTERS")
        apply_solid_fill(s, C_DARK_MAROON, border_color=C_MAROON, border_width=Pt(1.5))
    elif s.name == 'Text 47':
        set_tf_color(s.text_frame, C_WHITE)

prs.save('SIH26163_ARGUS_Idea_PPT_Color.pptx')
prs.save('sources/SIH26163_ARGUS_Idea_PPT_Color.pptx')
print("Complete Reference Image Color Palette successfully applied to all 6 slides!")
