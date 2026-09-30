import pptx

prs = pptx.Presentation('sources/SIH2026-IDEA-Presentation-Format.pptx')
print(f"Total slides: {len(prs.slides)}, width: {prs.slide_width.inches}, height: {prs.slide_height.inches}")

for idx in range(min(6, len(prs.slides))):
    slide = prs.slides[idx]
    print(f"\n==================== SLIDE {idx+1} ====================")
    for s_idx, shape in enumerate(slide.shapes):
        t = ""
        if shape.has_text_frame:
            t = " | ".join([p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()])
        print(f"[{s_idx}] Name: {shape.name} | Type: {shape.shape_type} | Pos: ({shape.left.inches:.2f}, {shape.top.inches:.2f}) | Size: ({shape.width.inches:.2f}x{shape.height.inches:.2f})")
        if t:
            print(f"    Text: {t[:120]}")
